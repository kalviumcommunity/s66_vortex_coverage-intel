"""Document loading and text extraction — FR-01, FR-02.

Spec: specs/04-ingestion.md
"""

from __future__ import annotations

import hashlib
import re
from pathlib import Path
from typing import Literal, Protocol

from docx import Document
from pydantic import BaseModel
from pypdf import PdfReader

from coverage_intel.errors import CorpusValidationError, ExtractionError, UnsupportedDocumentType

SUPPORTED_TYPES: frozenset[str] = frozenset({"pdf", "docx", "txt"})


class Block(BaseModel):
    """A contiguous span of extracted text with its structural provenance.

    Blocks — not one flat string — are what make section- and page-aware
    chunking possible (FR-02, FR-04).
    """

    text: str
    section: str | None = None
    page: int | None = None
    char_start: int | None = None
    char_end: int | None = None


class ExtractedDocument(BaseModel):
    """Result of extracting one source document."""

    document_id: str
    file_type: Literal["pdf", "docx", "txt"]
    blocks: list[Block]
    char_count: int


class DocumentLoader(Protocol):
    """Provider-agnostic loader interface."""

    def load(self, path: Path) -> ExtractedDocument: ...


def load(path: Path) -> ExtractedDocument:
    """Load and extract text from a PDF, DOCX, or TXT file."""
    suffix = path.suffix.lower().lstrip(".")
    if suffix not in SUPPORTED_TYPES:
        raise UnsupportedDocumentType(f"Unsupported document type: {path.suffix or '<none>'}")
    if not path.is_file():
        raise CorpusValidationError(f"Document does not exist: {path}")

    content = path.read_bytes()
    if not content:
        raise CorpusValidationError(f"Document is empty: {path.name}")
    document_id = f"{_slug(path.stem)}-{hashlib.sha256(content).hexdigest()[:8]}"

    if suffix == "pdf":
        blocks = _load_pdf(path)
    elif suffix == "docx":
        blocks = _load_docx(path)
    else:
        blocks = _load_txt(content)
    blocks = [block for block in blocks if block.text.strip()]
    char_count = sum(len(block.text) for block in blocks)
    if char_count == 0:
        raise ExtractionError(f"No usable text extracted from {path.name}")
    return ExtractedDocument(
        document_id=document_id, file_type=suffix, blocks=blocks, char_count=char_count
    )


def _slug(value: str) -> str:
    """Create a stable, readable identifier prefix."""
    slug = re.sub(r"[^a-z0-9]+", "-", value.lower()).strip("-")
    return slug or "document"


def _load_pdf(path: Path) -> list[Block]:
    return [
        Block(text=page.extract_text() or "", page=page_number)
        for page_number, page in enumerate(PdfReader(path).pages, start=1)
    ]


def _load_docx(path: Path) -> list[Block]:
    blocks: list[Block] = []
    section: str | None = None
    for paragraph in Document(path).paragraphs:
        text = paragraph.text.strip()
        if paragraph.style.name.startswith("Heading") and text:
            section = text
        if text:
            blocks.append(Block(text=text, section=section))
    return blocks


def _load_txt(content: bytes) -> list[Block]:
    text = _decode(content)
    blocks: list[Block] = []
    section: str | None = None
    char_offset = 0
    section_pattern = re.compile(r"^section\s+\S+$", re.IGNORECASE)
    uppercase_heading_pattern = re.compile(r"^[A-Z][A-Z\s\d&/-]{3,}$")
    for line in text.splitlines(keepends=True):
        value = line.strip()
        if value and (section_pattern.match(value) or uppercase_heading_pattern.match(value)):
            section = value
        if value:
            start = text.find(value, char_offset)
            blocks.append(
                Block(
                    text=value,
                    section=section,
                    page=1,
                    char_start=start,
                    char_end=start + len(value),
                )
            )
            char_offset = start + len(value)
    return blocks


def _decode(content: bytes) -> str:
    """Decode common document encodings without silently dropping bytes."""
    for encoding in ("utf-8-sig", "utf-16", "latin-1"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ExtractionError("Unable to decode text document")
